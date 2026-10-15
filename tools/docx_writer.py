"""
Write structured blocks (see tools/blocks.py) to Word, and read the document
back so Claude can verify what was written.

Reuses the styling, table, caption and COM field-update code in word_tool so
output matches documents produced by the Discord workflow.
"""

import os
import re
import shutil
import time

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt

from tools import word_tool as wt
from tools.blocks import inline_runs, latex_to_omml, unicode_fallback
from tools.parser import Block

_M_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def append_blocks(blocks: list[dict], word_file_path: str) -> str:
    styles = wt._load_styles()
    doc = wt._open_document(word_file_path, styles)
    needs_com = False

    for b in blocks:
        t = b["type"]
        if t in ("chapter", "heading2", "heading3", "body", "list_item"):
            text = b["text"]
            if t in ("heading2", "heading3"):
                # Word numbers headings itself; a typed "12.5 " would show twice.
                text = re.sub(r"^\s*\d+(\.\d+)+\s+", "", text)
            _write_text(doc, t, text, styles)
        elif t == "equation":
            _spacer(doc)
            _write_equation(doc, b, styles)
            _spacer(doc)
            needs_com = True
        elif t == "table":
            # Blank line before the table and after its caption, none in between.
            _spacer(doc)
            wt._write_table_block(doc, Block("table", table_rows=b["rows"]), styles, blank_after=False)
            wt._write_word_caption(doc, "Table", wt._strip_caption_prefix(b.get("caption", ""), "Table"))
            _spacer(doc)
            needs_com = True
        elif t == "figure":
            fig = Block("figure")
            fig.figure_number = b["number"]
            fig.figure_image_path = b.get("image", "")
            _spacer(doc)
            wt._write_figure_block(doc, fig, styles)
            wt._write_word_caption(doc, "Figure", wt._strip_caption_prefix(b.get("caption", ""), "Figure"))
            _spacer(doc)
            needs_com = True

    doc.save(word_file_path)
    if needs_com:
        wt._update_word_fields(word_file_path)
    return f"wrote {len(blocks)} block(s)"


def _is_blank(el) -> bool:
    """True for an empty paragraph: no text, no math, no image."""
    if el.tag != f"{_W_NS}p":
        return False
    if el.findall(f".//{_M_NS}oMath") or el.findall(f".//{_W_NS}drawing"):
        return False
    return not "".join(t.text or "" for t in el.iter(f"{_W_NS}t")).strip()


def _spacer(doc):
    """Add one blank line, unless the document already ends with one."""
    body = doc.element.body
    last = [el for el in body if el.tag != f"{_W_NS}sectPr"]
    if last and _is_blank(last[-1]):
        return
    doc.add_paragraph()


def _write_text(doc, block_type: str, text: str, styles: dict):
    cfg = styles.get(wt.STYLE_KEY_MAP.get(block_type, "body"), {})
    # Every chapter after the first starts on a new page (the template already
    # breaks the page before chapter 1).
    new_page = block_type == "chapter" and any(
        p.style is not None and p.style.name == "Heading 1" and p.text.strip()
        for p in doc.paragraphs
    )
    try:
        para = doc.add_paragraph(style=wt.WORD_STYLE_MAP.get(block_type, "Normal"))
    except KeyError:
        para = doc.add_paragraph()
    wt._apply_paragraph_format(para, cfg)
    if new_page:
        para.paragraph_format.page_break_before = True

    if block_type == "list_item":
        text = "• " + text
    for run in inline_runs(text):
        if run["math"]:
            omml = latex_to_omml(run["text"])
            if omml is not None:
                para._p.append(omml)
                continue
            r = para.add_run(unicode_fallback(run["text"]))
            wt._apply_run_format(r, cfg, italic=True)
            continue
        r = para.add_run(run["text"])
        wt._apply_run_format(r, cfg, bold=run["bold"], italic=run["italic"])


def _write_equation(doc, b: dict, styles: dict):
    cfg = styles.get("body", {})
    omml = latex_to_omml(b["latex"])
    para = doc.add_paragraph()
    label = (b.get("label") or "").strip()

    if label:
        sec = doc.sections[-1]
        width = sec.page_width - sec.left_margin - sec.right_margin
        stops = para.paragraph_format.tab_stops
        stops.add_tab_stop(int(width / 2), WD_TAB_ALIGNMENT.CENTER)
        stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        para.add_run("\t")
    else:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if omml is not None:
        para._p.append(omml)
    else:
        r = para.add_run(unicode_fallback(b["latex"]))
        r.italic = True

    if label:
        r = para.add_run(f"\t({label})")
        r.font.name = cfg.get("font", "Garamond")
        r.font.size = Pt(cfg.get("size_pt", 11))


# ── Read-back ─────────────────────────────────────────────────────────────────

def inspect(word_file_path: str, tail: int = 25) -> dict:
    """Summarise the document: counts, headings, and the last *tail* paragraphs."""
    if not os.path.exists(word_file_path):
        return {"exists": False}
    doc = Document(word_file_path)
    body = doc.element.body

    instr = " ".join(el.text or "" for el in body.iter(f"{_W_NS}instrText"))
    fig_captions = len(re.findall(r"SEQ Figure", instr))
    table_captions = len(re.findall(r"SEQ Table", instr))
    display_eqs = 0
    inline_eqs = 0
    for p in doc.paragraphs:
        maths = p._p.findall(f"{_M_NS}oMath") + p._p.findall(f".//{_M_NS}oMathPara")
        if not maths:
            continue
        has_text = any((r.text or "").strip() and not re.fullmatch(r"\(\S+\)", (r.text or "").strip())
                       for r in p.runs)
        if has_text:
            inline_eqs += len(maths)
        else:
            display_eqs += len(maths)

    headings = []
    for p in doc.paragraphs:
        name = p.style.name if p.style is not None else ""
        if name.startswith("Heading") and p.text.strip():
            headings.append(f"{name}: {p.text.strip()[:100]}")

    paras = [p for p in doc.paragraphs if p.text.strip() or p._p.findall(f".//{_M_NS}oMath")]
    tail_out = []
    for p in paras[-tail:]:
        name = p.style.name if p.style is not None else ""
        txt = p.text.strip()
        if p._p.findall(f".//{_M_NS}oMath"):
            txt += "  [contains equation]"
        tail_out.append(f"[{name}] {txt[:300]}")

    return {
        "exists": True,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "figure_captions": fig_captions,
        "table_captions": table_captions,
        "display_equations": display_eqs,
        "inline_equations": inline_eqs,
        "next_figure_number": fig_captions + 1,
        "next_table_number": table_captions + 1,
        "headings": headings,
        "tail": tail_out,
    }


def backup(word_file_path: str, backup_dir: str, keep: int = 15) -> str | None:
    if not os.path.exists(word_file_path):
        return None
    os.makedirs(backup_dir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(word_file_path))[0][:40]
    dest = os.path.join(backup_dir, f"{time.strftime('%Y%m%d-%H%M%S')}_{stem}.docx")
    shutil.copy2(word_file_path, dest)
    olds = sorted(f for f in os.listdir(backup_dir) if f.endswith(".docx"))
    for f in olds[:-keep]:
        os.remove(os.path.join(backup_dir, f))
    return dest
